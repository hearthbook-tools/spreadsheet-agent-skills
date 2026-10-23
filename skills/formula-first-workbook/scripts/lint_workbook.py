#!/usr/bin/env python3
"""Lint an .xlsx for formula mistakes. Usage: python lint_workbook.py FILE.xlsx [--json]. Requires openpyxl."""
import json
import re
import sys
from collections import Counter, defaultdict

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

REF = re.compile(r"(?<![A-Za-z_\d.!])(\$?)([A-Z]{1,3})(\$?)(\d{1,7})(?![\d(A-Za-z_])")
SHEET_REF = re.compile(r"(?:'([^']+)'|([A-Za-z_][\w.]*))!")
SUM_RANGE = re.compile(r"SUM\(\s*\$?([A-Z]{1,3})\$?(\d+)\s*:\s*\$?([A-Z]{1,3})\$?(\d+)\s*\)", re.I)
MAGIC = re.compile(r"(?<![A-Za-z_\d$.:])(\d+\.\d+|\d{2,})(?![\d:A-Za-z_(])")
OK_CONSTANTS = {"0", "1", "2", "7", "12", "14", "24", "30", "52", "60", "90", "100", "365", "366", "1000"}
AGGREGATE = re.compile(r"^=\s*(SUM|SUBTOTAL|AVERAGE|MIN|MAX|COUNT|COUNTA)\(", re.I)
ERRORS = ("#DIV/0!", "#N/A", "#NAME?", "#NULL!", "#NUM!", "#REF!", "#VALUE!")


def strip_strings(f):
    return re.sub(r'"[^"]*"', '""', f)


def pattern(formula, row, col):
    """Relative R1C1-style signature so filled-down formulas compare equal."""
    def rep(m):
        cabs, c, rabs, r = m.group(1), m.group(2), m.group(3), int(m.group(4))
        ci = column_index_from_string(c)
        cs = f"C{ci}" if cabs else f"C[{ci - col}]"
        rs = f"R{r}" if rabs else f"R[{r - row}]"
        return rs + cs
    return REF.sub(rep, strip_strings(formula))


def lint(path):
    wb = openpyxl.load_workbook(path)
    try:
        vals = openpyxl.load_workbook(path, data_only=True)
    except Exception:
        vals = None
    names = {ws.title for ws in wb.worksheets}
    issues = []
    add = lambda lvl, ws, cell, msg: issues.append({"level": lvl, "sheet": ws, "cell": cell, "msg": msg})  # noqa: E731
    cached_any = False
    for ws in wb.worksheets:
        cols = defaultdict(list)  # col -> [(row, kind, pattern)]
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and v.startswith("="):
                    f = v
                    if "#REF!" in f:
                        add("ERROR", ws.title, c.coordinate, "Broken reference (#REF!) in formula")
                    for m in SHEET_REF.finditer(strip_strings(f)):
                        sn = m.group(1) or m.group(2)
                        if sn not in names:
                            add("ERROR", ws.title, c.coordinate, f"Refers to missing sheet '{sn}'")
                    body = strip_strings(f)
                    nums = [n for n in MAGIC.findall(REF.sub("", body)) if n not in OK_CONSTANTS]
                    if nums:
                        add("WARN", ws.title, c.coordinate,
                            f"Number(s) typed inside formula: {', '.join(nums[:3])} - move to an Inputs cell")
                    for m in SUM_RANGE.finditer(body):
                        c1, r1, c2, r2 = m.group(1), int(m.group(2)), m.group(3), int(m.group(4))
                        if c1 == c2 and column_index_from_string(c1) == c.column and r2 < c.row - 1:
                            gap = [ws.cell(row=r, column=c.column).value for r in range(r2 + 1, c.row)]
                            if any(isinstance(g, (int, float)) or (isinstance(g, str) and g.startswith("="))
                                   for g in gap):
                                add("ERROR", ws.title, c.coordinate,
                                    f"SUM stops at row {r2} but values continue to row {c.row - 1}")
                    kind = "agg" if AGGREGATE.match(f) else "f"   # totals are allowed to differ from the column
                    cols[c.column].append((c.row, kind, pattern(f, c.row, c.column)))
                    if vals is not None:
                        cv = vals[ws.title][c.coordinate].value
                        if cv is not None:
                            cached_any = True
                        if isinstance(cv, str) and cv.startswith(ERRORS):
                            add("ERROR", ws.title, c.coordinate, f"Formula returns {cv}")
                elif isinstance(v, (int, float)) and not isinstance(v, bool):
                    cols[c.column].append((c.row, "n", None))
        for col, cells in cols.items():
            cells.sort()
            # contiguous blocks
            blocks, cur = [], [cells[0]] if cells else []
            for prev, nxt in zip(cells, cells[1:]):
                if nxt[0] == prev[0] + 1:
                    cur.append(nxt)
                else:
                    blocks.append(cur)
                    cur = [nxt]
            if cur:
                blocks.append(cur)
            for b in blocks:
                fs = [x for x in b if x[1] == "f"]
                if len(fs) >= 3:
                    common, cnt = Counter(x[2] for x in fs).most_common(1)[0]
                    if cnt >= 3 and cnt / len(fs) >= 0.6:
                        for r, k, p in b:
                            addr = f"{get_column_letter(col)}{r}"
                            if k == "n" and b.index((r, k, p)) not in (0, len(b) - 1):
                                add("WARN", ws.title, addr, "Typed number inside a column of formulas "
                                                            "(formula overwritten?)")
                            elif k == "f" and p != common and r not in (b[-1][0],):
                                add("WARN", ws.title, addr, "Formula differs from the pattern used in this column")
    return issues, cached_any


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    issues, cached = lint(sys.argv[1])
    if "--json" in sys.argv:
        print(json.dumps({"issues": issues, "cached_values_present": cached}, indent=2))
        return
    order = {"ERROR": 0, "WARN": 1}
    issues.sort(key=lambda i: (order[i["level"]], i["sheet"], i["cell"]))
    e = sum(i["level"] == "ERROR" for i in issues)
    print(f"{sys.argv[1]}: {e} error(s), {len(issues) - e} warning(s)")
    if not cached:
        print("Note: no cached results found (file never recalculated in Excel/Sheets); "
              "formula results were not checked, only structure.")
    for i in issues:
        print(f"  {i['level']:5} {i['sheet']}!{i['cell']}: {i['msg']}")
    sys.exit(1 if e else 0)


if __name__ == "__main__":
    main()
