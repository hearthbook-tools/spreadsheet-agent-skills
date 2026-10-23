#!/usr/bin/env python3
"""Break-even & pricing workbook.
Usage: python break_even.py --price P --var-cost V [--fee-pct F] --fixed X [--target-profit T] [--target-margin M] --out OUT.xlsx"""
import argparse
import math

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

INPUT, CALC = PatternFill("solid", fgColor="FFF7D6"), PatternFill("solid", fgColor="EEF1F4")
HDR, HDR_FILL = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="2F3A45")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--price", type=float, required=True)
    ap.add_argument("--var-cost", type=float, required=True)
    ap.add_argument("--fee-pct", type=float, default=0.0)
    ap.add_argument("--fixed", type=float, required=True)
    ap.add_argument("--target-profit", type=float, default=0.0)
    ap.add_argument("--target-margin", type=float, default=0.5)
    ap.add_argument("--out", required=True)
    a = ap.parse_args()
    wb = Workbook()
    ws = wb.active
    ws.title = "Pricing"
    rows = [
        ("Inputs", None, None),
        ("Price per unit", a.price, "in"), ("Variable cost per unit", a.var_cost, "in"),
        ("Platform / payment fee (% of price)", a.fee_pct, "in"), ("Fixed costs per month", a.fixed, "in"),
        ("Target profit per month", a.target_profit, "in"), ("Target margin (on price)", a.target_margin, "in"),
        ("Results", None, None),
        ("Fee per unit", "=B2*B4", "calc"),
        ("Contribution per unit", "=B2-B3-B9", "calc"),
        ("Contribution margin %", "=IF(B2=0,0,B10/B2)", "calc"),
        ("Break-even units per month", '=IF(B10<=0,"Never: price below unit costs",ROUNDUP(B5/B10,0))', "calc"),
        ("Break-even revenue per month", '=IF(B10<=0,"n/a",B12*B2)', "calc"),
        ("Units for target profit", '=IF(B10<=0,"n/a",ROUNDUP((B5+B6)/B10,0))', "calc"),
        ("Price for target margin", '=IF(1-B4-B7<=0,"Margin + fee >= 100%",B3/(1-B4-B7))', "calc"),
    ]
    for i, (label, val, kind) in enumerate(rows, start=1):
        ws.cell(row=i + 0, column=1, value=label)
        if kind is None:
            ws.cell(row=i, column=1).font = Font(bold=True, size=12)
            continue
        c = ws.cell(row=i, column=2, value=val)
        c.fill = INPUT if kind == "in" else CALC
    # B2 price, B3 var cost, B4 fee %, B5 fixed, B6 target profit, B7 target margin,
    # B9 fee/unit, B10 contribution, B11 margin %, B12 break-even units, B13 BE revenue, B14 target units, B15 price
    for r, fmt in ((2, '"$"#,##0.00'), (3, '"$"#,##0.00'), (4, "0.0%"), (5, '"$"#,##0'), (6, '"$"#,##0'),
                   (7, "0.0%"), (9, '"$"#,##0.00'), (10, '"$"#,##0.00'), (11, "0.0%"), (12, "#,##0"),
                   (13, '"$"#,##0'), (14, "#,##0"), (15, '"$"#,##0.00')):
        ws.cell(row=r, column=2).number_format = fmt
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 40, 18
    # What-if: profit at prices (rows) x monthly units (cols)
    ws["D1"] = "What-if: monthly profit at different prices and volumes"
    ws["D1"].font = Font(bold=True, size=12)
    ws["D2"] = "Price \\ Units"
    ws["D2"].font, ws["D2"].fill = HDR, HDR_FILL
    contrib = a.price - a.var_cost - a.price * a.fee_pct
    base_units = math.ceil(a.fixed / contrib) if contrib > 0 else 10
    prices = [round(a.price * m, 2) for m in (0.7, 0.85, 1.0, 1.15, 1.3, 1.5)]
    units = [max(1, round(base_units * m)) for m in (0.5, 1.0, 1.5, 2.0, 3.0)]
    for j, u in enumerate(units):  # editable volumes (yellow)
        c = ws.cell(row=2, column=5 + j, value=u)
        c.fill, c.font = INPUT, Font(bold=True)
    for i, pr in enumerate(prices):
        r = 3 + i
        c = ws.cell(row=r, column=4, value=pr)  # editable prices (yellow)
        c.number_format, c.fill = '"$"#,##0.00', INPUT
        for j in range(len(units)):
            col = 5 + j
            L = ws.cell(row=2, column=col).column_letter
            ws.cell(row=r, column=col, value=f"=($D{r}*(1-$B$4)-$B$3)*{L}$2-$B$5").number_format = \
                '"$"#,##0;[Red]-"$"#,##0'
    ws.column_dimensions["D"].width = 14
    ws["A17"] = "Estimates from your inputs. The what-if table shows arithmetic, not expected demand."
    ws["A17"].font = Font(italic=True, size=9)
    wb.save(a.out)
    print(f"Wrote {a.out}.")


if __name__ == "__main__":
    main()
